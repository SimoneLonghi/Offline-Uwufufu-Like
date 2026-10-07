
import json
import sys
from collections import defaultdict
from pathlib import Path
from typing import Optional, Protocol, TypedDict, Union
from urllib.parse import urlencode
from zipfile import BadZipFile

from openpyxl import load_workbook, Workbook
from openpyxl.utils.exceptions import InvalidFileException

if getattr(sys, "frozen", False):
    APP_DIR = Path(sys.executable).resolve().parent
else:
    APP_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = APP_DIR / "output"

Participant = tuple[str, str]
Winner = tuple[int, Participant]


class TournamentData(TypedDict):
    round_name: str
    groups: dict[int, list[Participant]]


class TournamentController(Protocol):
    filename: str
    state_filename: str


class TournamentState(Protocol):
    controller: TournamentController
    current_girone_index: int
    vincitori: list[Winner]


ROUND_BY_CAPACITY = {
    1   : "VINCITORE",
    2   : "Finale",
    4   : "Semifinale",
    8   : "Quarti",
    16  : "Ottavi",
    32  : "Sedicesimi"
}


def round_by_participant_count(participant_count: int) -> str:
    """Restituisce il turno del tabellone minimo che contiene i partecipanti."""
    for capacity, round_name in sorted(ROUND_BY_CAPACITY.items()):
        if participant_count <= capacity:
            return round_name
    return "Turno"


def load_file_gironi_by_name(next_tournament_round: str) -> Optional[TournamentData]:
    return load_file_gironi(OUTPUT_DIR / next_tournament_round)


def get_stato_path(tournament_name: str) -> Path:
    return OUTPUT_DIR / f"saved status {tournament_name}.json"


def get_tournament_name_from_workbook(workbook_stem: str) -> str:
    """Ricava il nome base dai workbook generati per un turno successivo."""
    stem = workbook_stem
    round_names = {*ROUND_BY_CAPACITY.values(), "Turno"}
    for round_name in sorted(round_names, key=len, reverse=True):
        for separator in (" ", "_"):
            prefix = f"{round_name}{separator}"
            if stem.casefold().startswith(prefix.casefold()):
                return stem[len(prefix):]
    return stem


def load_file_gironi(filepath: Union[str, Path]) -> Optional[TournamentData]:
    gironi: dict[int, list[Participant]] = defaultdict(list)
    wb = None
    try:
        wb = load_workbook(filepath, data_only=True)
        ws = wb.active  # primo foglio

        # Legge intestazioni dalla prima riga
        headers = {
            str(cell.value).strip().upper(): idx
            for idx, cell in enumerate(ws[1], start=1)
            if cell.value is not None
        }
        required_headers = {"NOME", "LINK", "GIRONE"}
        missing_headers = required_headers - headers.keys()
        if missing_headers:
            raise ValueError(
                "Intestazioni mancanti: " + ", ".join(sorted(missing_headers))
            )

        participant_count = 0
                
        for row in ws.iter_rows(min_row=2, values_only=True):
            nome = str(row[headers["NOME"] - 1]).strip() if row[headers["NOME"] - 1] else ""
            link_val = str(row[headers["LINK"] - 1]).strip() if row[headers["LINK"] - 1] else ""
            link = "https://" + link_val if link_val and not link_val.lower().startswith("http") else link_val
            girone_raw = row[headers["GIRONE"] - 1]

            if not nome and not link and not girone_raw:
                continue
            if not nome or girone_raw is None:
                raise ValueError("Riga con nome o girone mancante")

            if not link:
                query = urlencode({"tbm": "isch", "q": nome})
                link = f"https://www.google.com/search?{query}"

            try:
                girone_val = int(girone_raw)
            except (TypeError, ValueError) as e:
                raise ValueError("Il girone deve essere un numero intero") from e
            if girone_val < 1:
                raise ValueError("Il numero del girone deve essere maggiore di zero")

            if nome and link:
                gironi[girone_val].append((nome, link))
                participant_count += 1

        if participant_count == 0:
            return None

        return {
            "round_name": round_by_participant_count(participant_count),
            "groups": dict(gironi),
        }
    except (OSError, BadZipFile, InvalidFileException, ValueError, KeyError, IndexError) as e:
        print(f"Errore durante il caricamento del file: {e}")
        return None
    finally:
        if wb is not None:
            wb.close()

def save_vincitori(tournament: TournamentState) -> str:
    naming_file = (
        round_by_participant_count(len(tournament.vincitori))
        if tournament.vincitori
        else "vincitori"
    )
    file_vincitori = f"{naming_file} {tournament.controller.filename}.xlsx"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT_DIR / file_vincitori
    wb = Workbook()
    ws = wb.active
    ws.title = "Vincitori"

    # Intestazioni
    ws.append(["OLD GIRONE", "NOME", "LINK", "GIRONE"])

    # Dati
    for girone_num, vincitore in tournament.vincitori:
        ws.append([girone_num, vincitore[0], vincitore[1], (girone_num + 1) // 2])

    try:
        wb.save(output_path)
    finally:
        wb.close()

    return file_vincitori

def save_stato(tournament: TournamentState) -> None:
    stato = {
        "current_girone_index": tournament.current_girone_index,
        "vincitori": [(g, v[0], v[1]) for g, v in tournament.vincitori],
    }
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    state_path = get_stato_path(tournament.controller.state_filename)
    with state_path.open("w", encoding="utf-8") as f:
        json.dump(stato, f, indent=4)

def load_stato(tournament: TournamentState) -> None:
    tournament.current_girone_index = 0
    tournament.vincitori = []
    path = get_stato_path(tournament.controller.state_filename)
    if not path.exists():
        return

    try:
        with path.open("r", encoding="utf-8") as f:
            stato = json.load(f)
        girone_index = stato.get("current_girone_index", 0)
        vincitori_raw = stato.get("vincitori", [])
        if not isinstance(girone_index, int) or girone_index < 0:
            raise ValueError("Indice del girone non valido")
        vincitori = []
        for entry in vincitori_raw:
            if not isinstance(entry, list) or len(entry) not in (3, 4):
                raise ValueError("Formato dei vincitori salvati non valido")
            group_number, name, link = entry[:3]
            vincitori.append((group_number, (name, link)))
    except (OSError, json.JSONDecodeError, KeyError, TypeError, ValueError) as e:
        print(f"Stato salvato non valido ({path}): {e}")
        return

    tournament.current_girone_index = girone_index
    tournament.vincitori = vincitori
